"""Render the private review for Feishu. Requires python-docx (authoring only)."""
from pathlib import Path
import re
from docx import Document
from docx.shared import Inches,Pt
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
root=Path(__file__).resolve().parents[1]
folder=root/'docs/release/internal'
doc=Document();section=doc.sections[0]
section.left_margin=section.right_margin=Inches(.55)
style=doc.styles['Normal'];style.font.name='Arial';style.font.size=Pt(10)
style.element.rPr.rFonts.set(qn('w:eastAsia'),'Microsoft YaHei')
lines=(folder/'GAP-REGISTER.md').read_text().splitlines();i=0
while i<len(lines):
    line=lines[i]
    if line.startswith('|'):
        rows=[]
        while i<len(lines) and lines[i].startswith('|'):
            parts=[p.strip() for p in lines[i].strip('|').split('|')]
            if not all(re.fullmatch('[-: ]+',p) for p in parts):rows.append(parts)
            i+=1
        table=doc.add_table(rows=0,cols=len(rows[0]));table.style='Table Grid'
        for row in rows:
            cells=table.add_row().cells
            for cell,value in zip(cells,row):
                cell.text=value.replace('`','').replace('**','')
                for p in cell.paragraphs:
                    for run in p.runs:run.font.size=Pt(8)
        repeat=OxmlElement('w:tblHeader');table.rows[0]._tr.get_or_add_trPr().append(repeat)
        continue
    if line.startswith('#'):
        count=len(line)-len(line.lstrip('#'));doc.add_heading(line[count:].strip(),min(count-1,3))
    elif line.strip():doc.add_paragraph(line.replace('`','').replace('**',''))
    i+=1
verification=root/'docs/release/VERIFICATION.md'
if verification.exists():
    doc.add_page_break();doc.add_heading('实际验证记录',0)
    for line in verification.read_text().splitlines():
        if line.strip():doc.add_paragraph(line.replace('`','').replace('**',''))
path=folder/'辽轨实训_差距与落地清单.docx';doc.save(path);print(path.name)
