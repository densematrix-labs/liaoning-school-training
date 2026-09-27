#!/usr/bin/env python3
"""将投标视频稿的性能段落更新为真实终端演示口径。"""

from __future__ import annotations

import argparse
import json
from pathlib import Path

from docx import Document


def set_paragraph(paragraph, text: str) -> None:
    if paragraph.runs:
        paragraph.runs[0].text = text
        for run in paragraph.runs[1:]:
            run.text = ""
    else:
        paragraph.add_run(text)


def replace_prefix(document: Document, prefix: str, text: str) -> None:
    for paragraph in document.paragraphs:
        if paragraph.text.strip().startswith(prefix):
            set_paragraph(paragraph, text)
            return
    raise RuntimeError(f"未找到待替换段落：{prefix}")


def replace_between(document: Document, start_prefix: str, end_prefix: str, values: list[str]) -> None:
    start = next(index for index, paragraph in enumerate(document.paragraphs) if paragraph.text.strip().startswith(start_prefix))
    end = next(index for index, paragraph in enumerate(document.paragraphs) if index > start and paragraph.text.strip().startswith(end_prefix))
    targets = document.paragraphs[start:end]
    if len(values) > len(targets):
        raise RuntimeError("性能段落空间不足，无法保持原文档结构")
    for paragraph, value in zip(targets, values):
        set_paragraph(paragraph, value)
    for paragraph in targets[len(values):]:
        set_paragraph(paragraph, "")


def replace_cell(cell, text: str) -> None:
    paragraph = cell.paragraphs[0]
    set_paragraph(paragraph, text)
    for extra in cell.paragraphs[1:]:
        set_paragraph(extra, "")


def main() -> int:
    parser = argparse.ArgumentParser()
    parser.add_argument("source", type=Path)
    parser.add_argument("result", type=Path)
    parser.add_argument("output", type=Path)
    args = parser.parse_args()
    metrics = json.loads(args.result.read_text(encoding="utf-8"))
    regular = metrics["groups"]["regular"]
    aggregate = metrics["groups"]["aggregate"]
    document = Document(args.source)

    replace_prefix(
        document,
        "推荐成片时长：",
        "推荐成片时长：9 分 05 秒\n演示地址： https://shixun.demo.densematrix.ai\n叙事顺序： 学生端 → 教师端 → 管理端 → 独立只读终端记录\n演示主线： 实训数据进入平台后，形成“成绩证据—能力评价—AI 诊断—教师干预—规则治理—同步审计”的完整闭环；性能佐证从产品界面切换到独立终端演示。",
    )
    performance_values = [
        "08:15–08:50｜性能与工程完整性佐证（独立终端）",
        "画面操作",
        "从管理端退出产品业务画面，切换到独立终端窗口；产品 UI 内不展示性能测试面板或测试控制。",
        "本机终端录制：先用 read -s PERF_TOKEN && export PERF_TOKEN 静默设置访问令牌，再执行 ./scripts/performance_demo.sh --duration 30。令牌不会出现在命令回显、原始日志或汇总结果中。",
        "无本地终端时，浏览器直接打开 https://shixun.demo.densematrix.ai/performance-terminal。该页面仅自动回放本轮脱敏终端记录，不是实时 Shell，不支持命令输入或任务触发。",
        "镜头先停留在“只读终端记录 / 实际测试时间 / 非实时 Shell”说明 3 秒，再依次停留常规查询指标、班级统计/能力汇总指标和总体结果各 4 秒；可使用暂停、重新播放或全屏。",
        "解说词",
        f"本轮在公网 Demo 上进行了 50 个并发虚拟用户的只读业务验证。测试实际持续 {metrics['actual_duration_seconds']:.3f} 秒，共完成 {metrics['total_requests']} 次 GET 请求，成功 {metrics['success']} 次、失败 {metrics['failed']} 次，成功率 {metrics['success_rate']:.3f}%，总体吞吐 {metrics['throughput_rps']:.3f} 次每秒。",
        f"其中，常规查询 P95 为 {regular['p95_seconds']:.3f} 秒，低于 3 秒目标；班级统计与能力汇总 P95 为 {aggregate['p95_seconds']:.3f} 秒，低于 10 秒目标。终端同时展示各组请求数、成功与失败、吞吐以及 P50、P95、P99 和阈值判断。",
        "以上为 Demo 工程验证，不替代正式验收。正式验收需由采购人与中标人共同确认服务器、网络、测试数据量、账号和标准业务脚本，并以原始日志、机器可读结果和汇总报告形成验收记录。",
        "画面数据口径",
        f"实际测试时间 {metrics['tested_at']}；{metrics['environment']}；{metrics['concurrency']} 并发；{metrics['actual_duration_seconds']:.3f} 秒。常规查询：{regular['requests']} 次、P50 {regular['p50_seconds']:.3f} 秒、P95 {regular['p95_seconds']:.3f} 秒、P99 {regular['p99_seconds']:.3f} 秒；班级统计/能力汇总：{aggregate['requests']} 次、P50 {aggregate['p50_seconds']:.3f} 秒、P95 {aggregate['p95_seconds']:.3f} 秒、P99 {aggregate['p99_seconds']:.3f} 秒。总体 {metrics['total_requests']} 次请求，成功 {metrics['success']}、失败 {metrics['failed']}、成功率 {metrics['success_rate']:.3f}%、吞吐 {metrics['throughput_rps']:.3f} 次/秒；本轮 Demo 工程阈值判断通过。",
        "录制控制：不得展示真实令牌、Cookie、账号口令、服务器路径、内部 IP、完整内部 URL、学生隐私或终端历史命令。回放页无产品导航、无命令输入框、无执行按钮，仅保留暂停、重新播放和全屏。",
    ]
    replace_between(document, "08:15–08:35｜性能与工程完整性佐证", "08:35–08:50｜片尾总结", performance_values)
    replace_prefix(document, "08:35–08:50｜片尾总结", "08:50–09:05｜片尾总结")

    timeline = document.tables[0]
    replace_cell(timeline.rows[4].cells[0], "05:30–08:50")
    replace_cell(timeline.rows[4].cells[2], "全校总览、规则与映射重算、千条同步与异常隔离、权限与运行状态；随后切换独立终端性能佐证")
    replace_cell(timeline.rows[5].cells[0], "08:50–09:05")

    coverage = document.tables[1]
    performance_row = next(row for row in coverage.rows if "8.1.2" in row.cells[0].text)
    replace_cell(performance_row.cells[2], "08:15–08:50")
    replace_cell(performance_row.cells[3], "独立终端或只读回放直链；50 并发；两类 P95；原始日志、机器结果、汇总报告及正式验收口径")

    checklist_replacements = {
        "50 并发测试采用最新一次复测结果": f"50 并发只读测试采用 {metrics['tested_at']} 公网复测结果：{metrics['total_requests']} 次请求、成功率 {metrics['success_rate']:.3f}%、常规查询 P95 {regular['p95_seconds']:.3f} 秒、班级统计/能力汇总 P95 {aggregate['p95_seconds']:.3f} 秒；视频数字须与脱敏原始日志、机器结果和汇总文本一致。",
        "全片避免出现桌面隐私": "全片避免出现桌面隐私、浏览器自动填充信息、密钥、Cookie、账号口令、服务器地址、内部 IP、终端历史命令和无关聊天通知；终端窗口只展示脚本脱敏输出。",
    }
    for prefix, value in checklist_replacements.items():
        replace_prefix(document, prefix, value)
    anchor = next(paragraph for paragraph in document.paragraphs if paragraph.text.strip().startswith("50 并发只读测试采用"))
    insert = anchor._p.addnext
    new_paragraph = document.add_paragraph()
    set_paragraph(new_paragraph, "录制前分别检查本机一键脚本和 https://shixun.demo.densematrix.ai/performance-terminal；回放页应明确“只读终端记录、实际测试时间、非实时 Shell”，且仅有暂停、重新播放和全屏。")
    anchor._p.addnext(new_paragraph._p)

    args.output.parent.mkdir(parents=True, exist_ok=True)
    document.save(args.output)
    reopened = Document(args.output)
    visible_text = "\n".join(paragraph.text for paragraph in reopened.paragraphs)
    visible_text += "\n" + "\n".join(cell.text for table in reopened.tables for row in table.rows for cell in row.cells)
    forbidden = ["测试结果页或报告截屏", "2026 年 9 月 4 日", "0.801 秒", "0.736 秒"]
    if any(item in visible_text for item in forbidden):
        raise RuntimeError("文档仍包含旧性能口径")
    if "performance-terminal" not in visible_text or f"{metrics['total_requests']} 次" not in visible_text:
        raise RuntimeError("文档缺少回放链接或本轮真实指标")
    print(f"generated={args.output}")
    print(f"paragraphs={len(reopened.paragraphs)} tables={len(reopened.tables)}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
