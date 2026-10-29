from fastapi import APIRouter, BackgroundTasks, Depends, HTTPException
from fastapi.responses import Response
from sqlalchemy.ext.asyncio import AsyncSession
from typing import List, Optional
import html

from app.database import get_db
from app.config import settings
from app.routers.auth import get_current_user
from app.services.report import ReportService, process_report_task
from app.schemas.auth import UserResponse
from app.schemas.report import DiagnosticReportResponse, GenerateReportRequest, ReportTaskResponse
from app.routers.permissions import require_student_access

router = APIRouter(prefix="/api/v1/reports", tags=["诊断报告"])


@router.get('/{report_id}/download.docx')
async def download_docx(report_id:str,current_user:UserResponse=Depends(get_current_user),db:AsyncSession=Depends(get_db)):
    import io,re
    from docx import Document
    from docx.shared import Pt
    from docx.oxml.ns import qn
    result=await ReportService(db).get_report(report_id)
    if not result:
        raise HTTPException(404,'报告不存在')
    await require_student_access(current_user,result.student_id,db)
    doc=Document();style=doc.styles['Normal'];style.font.name='宋体';style.font.size=Pt(11)
    style.element.rPr.rFonts.set(qn('w:eastAsia'),'宋体');doc.add_heading(result.title,0)
    lines=result.content.splitlines();i=0
    while i<len(lines):
        line=lines[i]
        if line.startswith('|') and i+1<len(lines) and re.match(r'^\|[\s:|\-]+\|$',lines[i+1]):
            cells=[x.strip() for x in line.strip('|').split('|')];table=doc.add_table(rows=1,cols=len(cells));table.style='Table Grid'
            for c,t in zip(table.rows[0].cells,cells):c.text=t
            i+=2
            while i<len(lines) and lines[i].startswith('|'):
                row=table.add_row().cells
                for c,t in zip(row,lines[i].strip('|').split('|')):c.text=t.strip()
                i+=1
            continue
        clean=re.sub(r'\*\*(.*?)\*\*',r'\1',line)
        if clean.startswith('#'):doc.add_heading(clean.lstrip('#').strip(),min(3,len(clean)-len(clean.lstrip('#'))))
        elif clean.strip():doc.add_paragraph(clean.lstrip('- ') if clean.startswith('- ') else clean)
        i+=1
    stream=io.BytesIO();doc.save(stream)
    return Response(stream.getvalue(),media_type='application/vnd.openxmlformats-officedocument.wordprocessingml.document',headers={'Content-Disposition':f'attachment; filename="{result.report_reference or "diagnostic-report"}.docx"'})


@router.post("/generate", response_model=ReportTaskResponse)
async def generate_report(
    request: GenerateReportRequest,
    background_tasks: BackgroundTasks,
    current_user: UserResponse = Depends(get_current_user),
    db: AsyncSession = Depends(get_db),
):
    """生成诊断报告"""
    await require_student_access(current_user, request.student_id, db)
    
    service = ReportService(db)
    
    try:
        task = await service.create_task(
            student_id=request.student_id,
            report_type=request.report_type,
            score_id=request.score_id,
            created_by=current_user.id,
            date_from=request.date_from,
            date_to=request.date_to,
        )
        if not settings.RELEASE_MODE:
            background_tasks.add_task(process_report_task, task.id)
        return task
    except ValueError as e:
        raise HTTPException(status_code=400, detail=str(e))
    except RuntimeError as e:
        raise HTTPException(status_code=502, detail=str(e))


@router.get("/tasks/{task_id}", response_model=ReportTaskResponse)
async def get_report_task(
    task_id: str,
    current_user: UserResponse = Depends(get_current_user),
    db: AsyncSession = Depends(get_db),
):
    task = await ReportService(db).get_task(task_id)
    if not task:
        raise HTTPException(status_code=404, detail="报告任务不存在")
    await require_student_access(current_user, task.student_id, db)
    return task


@router.get("/", response_model=List[DiagnosticReportResponse])
async def get_my_reports(
    limit: int = 10,
    current_user: UserResponse = Depends(get_current_user),
    db: AsyncSession = Depends(get_db),
):
    """获取当前学生的诊断报告列表"""
    if current_user.role != "student" or not current_user.student_id:
        raise HTTPException(status_code=403, detail="仅学生可访问")
    
    service = ReportService(db)
    return await service.get_student_reports(current_user.student_id, limit)


@router.get("/student/{student_id}", response_model=List[DiagnosticReportResponse])
async def get_student_reports(
    student_id: str,
    limit: int = 10,
    report_type: Optional[str] = None,
    score_id: Optional[str] = None,
    current_user: UserResponse = Depends(get_current_user),
    db: AsyncSession = Depends(get_db),
):
    """获取指定学生的诊断报告列表（教师/管理员）"""
    await require_student_access(current_user, student_id, db)
    
    service = ReportService(db)
    try:
        return await service.get_student_reports(
            student_id,
            limit,
            report_type=report_type,
            score_id=score_id,
        )
    except ValueError as exc:
        raise HTTPException(status_code=400, detail=str(exc)) from exc


@router.get("/{report_id}", response_model=DiagnosticReportResponse)
async def get_report(
    report_id: str,
    current_user: UserResponse = Depends(get_current_user),
    db: AsyncSession = Depends(get_db),
):
    """获取报告详情"""
    service = ReportService(db)
    result = await service.get_report(report_id)
    
    if not result:
        raise HTTPException(status_code=404, detail="报告不存在")
    
    await require_student_access(current_user, result.student_id, db)
    
    return result


@router.get("/{report_id}/download.doc")
async def download_report(
    report_id: str,
    current_user: UserResponse = Depends(get_current_user),
    db: AsyncSession = Depends(get_db),
):
    result = await ReportService(db).get_report(report_id)
    if not result:
        raise HTTPException(status_code=404, detail="报告不存在")
    await require_student_access(current_user, result.student_id, db)
    body = f"""<html><head><meta charset=\"utf-8\"><title>{html.escape(result.title)}</title></head><body><h1>{html.escape(result.title)}</h1><pre style=\"white-space:pre-wrap;font-family:SimSun,serif\">{html.escape(result.content)}</pre></body></html>"""
    return Response(
        body.encode("utf-8"),
        media_type="application/msword",
        headers={"Content-Disposition": f'attachment; filename="{result.report_reference or "diagnostic-report"}.doc"'},
    )
